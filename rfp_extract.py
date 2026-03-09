# rfp_extract.py — multi-file extractor (DOCX + XLSX)
# - Strikte requirement-poort + legal/adres-filter
# - Hybride score (rule + semantic)
# - Governance gating (cues + min similarity)
# - Uitleg fail-safe (chat timeout / fallback)
from __future__ import annotations

import argparse, sys, glob
from pathlib import Path
import pandas as pd
import docx

from utils.helpers import (
    normalize_ws,
    rule_requirement_score,
    is_probably_requirement_text,
    looks_like_legal_or_address,
    classify_requirement,
    detect_document_role,
)
from ai.semantic_classifier import SemanticRequirementClassifier
from ai.semantic_governance import SemanticGovernanceMapper

OUTPUT_COLS = [
    "id", "text", "type", "source_file", "source_role",
    "cobit", "itil", "iso27001",
    "governance_explanation", "governance_confidence"
]
ALLOWED_EXTS = {".docx", ".xlsx"}

def collect_input_files(inputs: list[str]) -> list[Path]:
    files : list[Path] = []
    for arg in inputs:
        p = Path(arg)
        if p.is_dir():
            files += list(p.glob("*.docx")) + list(p.glob("*.xlsx"))
        else:
            matches = glob.glob(str(p))
            if not matches and p.exists():
                matches = [str(p)]
            for m in matches:
                mp = Path(m)
                if mp.suffix.lower() in ALLOWED_EXTS and mp.exists():
                    files.append(mp)
    uniq, seen = [], set()
    for f in files:
        if f not in seen:
            uniq.append(f)
            seen.add(f)
    return uniq

def parse_docx_file(path: Path, clf: SemanticRequirementClassifier, mapper: SemanticGovernanceMapper):
    out = []
    try:
        document = docx.Document(str(path))
    except Exception as e:
        print(f"[WARN] Cannot open DOCX: {path} ({e})", file=sys.stderr)
        return out

    # Document role
    header_text = []
    for p in document.paragraphs[:15]:
        t = normalize_ws(p.text)
        if t:
            header_text.append(t)
    role = detect_document_role("\n".join(header_text), path.name)

    # PRICING/LEGAL → skip
    if role in {"PRICING", "LEGAL"}:
        print(f"[INFO] Skipping requirement extraction for {path.name} (role={role})")
        return out

    for p in document.paragraphs:
        txt = normalize_ws(p.text)
        if not txt:
            continue
        # absolute legal/adres skip
        if looks_like_legal_or_address(txt):
            continue
        # strikte poort
        if not is_probably_requirement_text(txt):
            continue

        # hybride score
        rule_s = rule_requirement_score(txt)
        sem_s  = clf.score(txt)
        final  = 0.5 * rule_s + 0.5 * sem_s

        # strengere drempel voor BR uit RFP_MAIN
        rtype = classify_requirement(txt, role_hint=role)
        if role == "RFP_MAIN" and rtype == "BR":
            if final < 0.65:
                continue
        else:
            if final < 0.55:
                continue

        gov = mapper.map_text(txt)
        out.append({
            "id": mapper.hash_id(txt),
            "text": txt,
            "type": rtype,
            "source_file": path.name,
            "source_role": role,
            "cobit": ", ".join(gov.cobit),
            "itil": ", ".join(gov.itil),
            "iso27001": ", ".join(gov.iso),
            "governance_explanation": gov.explanation,
            "governance_confidence": gov.confidence,
        })
    return out

def _pick_text_column(columns):
    cand = {c.lower().strip(): c for c in columns}
    priority = ["requirement", "requirements", "req", "requirement_text",
                "description", "omschrijving", "beschrijving", "text"]
    for key in priority:
        for k, v in cand.items():
            if key == k:
                return v
    for k, v in cand.items():
        if any(x in k for x in ["require", "desc", "text", "eis"]):
            return v
    return columns[0] if columns else None

def parse_xlsx_file(path: Path, clf: SemanticRequirementClassifier, mapper: SemanticGovernanceMapper):
    out = []
    import openpyxl  # engine
    try:
        xls = pd.ExcelFile(str(path), engine="openpyxl")
    except Exception as e:
        print(f"[WARN] Cannot open XLSX: {path} ({e})", file=sys.stderr)
        return out

    role = "GENERIC"
    try:
        df0 = pd.read_excel(xls, sheet_name=xls.sheet_names[0], engine="openpyxl", nrows=100)
        sample_text = " ".join([str(x) for x in df0.astype(str).values.flatten().tolist()[:200]])
        role = detect_document_role(sample_text, path.name)
    except Exception:
        pass

    if role in {"PRICING", "LEGAL"}:
        print(f"[INFO] Skipping requirement extraction for {path.name} (role={role})")
        return out

    for sheet in xls.sheet_names:
        try:
            df = pd.read_excel(xls, sheet_name=sheet, engine="openpyxl")
        except Exception as e:
            print(f"[WARN] Cannot read sheet {path.name}:{sheet} ({e})", file=sys.stderr)
            continue
        if df.empty:
            continue
        text_col = _pick_text_column(list(df.columns))
        if not text_col:
            continue

        for _, row in df.iterrows():
            txt = normalize_ws(str(row.get(text_col, "")))
            if not txt:
                continue
            if looks_like_legal_or_address(txt):
                continue
            if not is_probably_requirement_text(txt):
                continue

            rule_s = rule_requirement_score(txt)
            sem_s  = clf.score(txt)
            final  = 0.5 * rule_s + 0.5 * sem_s

            rtype = classify_requirement(txt, role_hint=role)
            if role == "RFP_MAIN" and rtype == "BR":
                if final < 0.65:
                    continue
            else:
                if final < 0.55:
                    continue

            gov = mapper.map_text(txt)
            out.append({
                "id": mapper.hash_id(txt),
                "text": txt,
                "type": rtype,
                "source_file": f"{path.name}:{sheet}",
                "source_role": role,
                "cobit": ", ".join(gov.cobit),
                "itil": ", ".join(gov.itil),
                "iso27001": ", ".join(gov.iso),
                "governance_explanation": gov.explanation,
                "governance_confidence": gov.confidence,
            })
    return out

def dedupe(records: list[dict]) -> list[dict]:
    seen, out = set(), []
    for r in records:
        key = normalize_ws(r["text"]).lower()
        if key in seen:
            continue
        seen.add(key)
        out.append(r)
    return out

def export_xlsx(records: list[dict], out_path: Path):
    df = pd.DataFrame(records)
    for c in OUTPUT_COLS:
        if c not in df.columns:
            df[c] = ""
    df = df[OUTPUT_COLS]

    tabs = {
        "BR": df[df["type"] == "BR"],
        "FR": df[df["type"] == "FR"],
        "NFR": df[df["type"] == "NFR"],
        "RISKS": df[df["type"] == "RISK"],
        "SCOPE_IN": df[df["type"] == "SCOPE_IN"],
        "SCOPE_OUT": df[df["type"] == "SCOPE_OUT"],
        "SCOPE_OTHER": df[df["type"] == "SCOPE_OTHER"],
        "ALL": df,
    }
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        for name, sub in tabs.items():
            sub.to_excel(writer, sheet_name=name, index=False)

def main():
    parser = argparse.ArgumentParser(description="RFP multi-file extractor (hybrid filter + governance gating)")
    parser.add_argument("inputs", nargs="+", help="Files/patterns or a directory (e.g., examples/ or 'examples/*.docx' 'examples/*.xlsx')")
    parser.add_argument("--xlsx", default="out.xlsx")
    parser.add_argument("--ai", default="on", choices=["on", "off"])
    parser.add_argument("--embed-model", default=None)
    parser.add_argument("--chat-model",  default=None)
    args = parser.parse_args()

    files = collect_input_files(args.inputs)
    if not files:
        print("[ERR] No input files found.")
        sys.exit(2)

    print("[INFO] Files:")
    for f in files:
        print(f"  - {f}")

    clf    = SemanticRequirementClassifier(embed_model=args.embed_model)
    mapper = SemanticGovernanceMapper(embed_model=args.embed_model, chat_model=args.chat_model)

    all_recs : list[dict] = []
    for f in files:
        if f.suffix.lower() == ".docx":
            all_recs += parse_docx_file(f, clf, mapper)
        elif f.suffix.lower() == ".xlsx":
            all_recs += parse_xlsx_file(f, clf, mapper)

    merged = dedupe(all_recs)
    export_xlsx(merged, Path(args.xlsx))
    print(f"[OK] Wrote {args.xlsx} (records={len(merged)})")

if __name__ == "__main__":
    main()